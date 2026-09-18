from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'app/src/main/assets/templates'; OUT.mkdir(exist_ok=True)
FONT='TH Sarabun New'; GARUDA=ROOT/'app/src/main/assets/garuda.png'
def font(r,b=False):
 r.font.name=FONT;r.font.size=Pt(16);r.bold=b
 rPr=r._r.get_or_add_rPr();rf=OxmlElement('w:rFonts')
 for k in ('ascii','hAnsi','cs','eastAsia'):rf.set(qn('w:'+k),FONT)
 rPr.append(rf)
def base(d):
 s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.5);s.bottom_margin=Cm(2);s.left_margin=Cm(3);s.right_margin=Cm(2)
 st=d.styles['Normal'];st.font.name=FONT;st.font.size=Pt(16);st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(0);st.paragraph_format.line_spacing=1.0
def p(d,t='',a=None,b=False,first=0,left=0):
 x=d.add_paragraph();x.paragraph_format.space_before=Pt(0);x.paragraph_format.space_after=Pt(0);x.paragraph_format.line_spacing=1.0
 if first:x.paragraph_format.first_line_indent=Cm(first)
 if left:x.paragraph_format.left_indent=Cm(left)
 if a is not None:x.alignment=a
 r=x.add_run(t);font(r,b);return x
def tabp(d,a,b,tab=10):
 x=d.add_paragraph();x.paragraph_format.space_before=Pt(0);x.paragraph_format.space_after=Pt(0);x.paragraph_format.line_spacing=1.0;x.paragraph_format.tab_stops.add_tab_stop(Cm(tab))
 for i,t in enumerate((a,b)):
  if i:x.add_run('\t')
  r=x.add_run(t);font(r)
def garuda(d):
 x=d.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(0);x.add_run().add_picture(str(GARUDA),width=Cm(2.5))
# external
D=Document();base(D);garuda(D);tabp(D,'ที่ {{DOC_NO}}','{{SENDER_ORG}}');tabp(D,'','{{SENDER_ADDR}}');tabp(D,'','{{DATE}}');p(D,'เรื่อง   {{SUBJECT}}');p(D,'เรียน   {{RECIPIENTS}}',left=.5);p(D,'สิ่งที่ส่งมาด้วย   {{ATTACHMENT}}',left=.5);p(D,'{{BODY_1}}',first=2.5);p(D,'{{BODY_2}}',first=2.5);p(D,'{{BODY_3}}',first=2.5);p(D,'ขอแสดงความนับถือ',a=WD_ALIGN_PARAGRAPH.CENTER);p(D,'(ลงชื่อ) {{SIGNER_NAME}}',a=WD_ALIGN_PARAGRAPH.CENTER);p(D,'{{SIGNER_TITLE}}',a=WD_ALIGN_PARAGRAPH.CENTER);p(D,'{{CONTACT_DEPT}}');p(D,'โทร. {{PHONE}}');p(D,'ผู้ประสานงาน {{CONTACT}}');D.save(OUT/'external_template.docx')
# memo
D=Document();base(D);garuda(D)
p(D,'บันทึกข้อความ',a=WD_ALIGN_PARAGRAPH.CENTER,b=True);p(D,'ส่วนราชการ   {{SENDER_ORG}}');tabp(D,'ที่ {{DOC_NO}}','วันที่ {{DATE}}',9);p(D,'เรื่อง   {{SUBJECT}}');p(D,'เรียน   {{RECIPIENTS}}');p(D,'๑. เรื่องเดิม',b=True);p(D,'{{OLD_MATTER}}',first=2.5);p(D,'๒. ข้อเท็จจริง',b=True);p(D,'{{FACTS}}',first=2.5);p(D,'๓. ข้อพิจารณา',b=True);p(D,'{{CONSIDERATION}}',first=2.5);p(D,'{{CLOSING}}',first=2.5);p(D,'(ลงชื่อ) {{SIGNER_NAME}}',a=WD_ALIGN_PARAGRAPH.CENTER);p(D,'{{SIGNER_TITLE}}',a=WD_ALIGN_PARAGRAPH.CENTER);D.save(OUT/'memo_template.docx')
# command
D=Document();base(D);garuda(D)
p(D,'{{ORG_NAME}}',a=WD_ALIGN_PARAGRAPH.CENTER,b=True);p(D,'ที่ {{DOC_NO}}',a=WD_ALIGN_PARAGRAPH.CENTER);p(D,'เรื่อง  {{SUBJECT}}',a=WD_ALIGN_PARAGRAPH.CENTER);p(D,'----------------------------------------------',a=WD_ALIGN_PARAGRAPH.CENTER);p(D,'{{PREAMBLE}}',first=2.5);p(D,'{{LIST_1}}');p(D,'{{LIST_2}}');p(D,'{{LIST_3}}');p(D,'{{DUTIES}}',first=2.5);p(D,'{{EFFECTIVE}}',first=2.5);p(D,'สั่ง ณ วันที่ {{DATE}}',a=WD_ALIGN_PARAGRAPH.CENTER);p(D,'(ลงชื่อ) {{SIGNER_NAME}}',a=WD_ALIGN_PARAGRAPH.CENTER);p(D,'{{SIGNER_TITLE}}',a=WD_ALIGN_PARAGRAPH.CENTER);D.save(OUT/'command_template.docx')
